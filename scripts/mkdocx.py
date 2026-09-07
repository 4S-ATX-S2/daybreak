# -*- coding: utf-8 -*-
"""Native DOCX builder for the Daybreak Brief.
RULE: a Word paragraph == an HTML *block*. Every inline descendant (<b>, <em>,
<span>, bare text) becomes a RUN in the SAME paragraph. CSS flex/grid rows carry
no source whitespace, so separators are inserted explicitly.
"""
import sys, re, os
from bs4 import BeautifulSoup, NavigableString, Tag
from docx import Document
from docx.shared import Pt, Inches, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

GOLD = RGBColor(0x94, 0x80, 0x3a); INK = RGBColor(0x26, 0x25, 0x1f)
SOFT = RGBColor(0x54, 0x52, 0x4a); RED = RGBColor(0xa4, 0x33, 0x1f)
BAR  = RGBColor(0x21, 0x1f, 0x18)
CHIPBG = {'c-calm':'7d8a4e','c-attentive':'a08b3f','c-elevated':'b0592b','c-high':'7c1f1f','c-severe':'511414'}
BADGEBG = {'b-open':'a08b3f','b-monitoring':'b0592b','b-new':'211f18','b-closed':'8a8578'}

INLINE = {'b','strong','i','em','span','a','sup','sub','u','small','code','br'}
SCALE = 1.0   # set by build()

def shade(cell, hexcolor):
    tcPr = cell._tc.get_or_add_tcPr()
    sh = OxmlElement('w:shd'); sh.set(qn('w:val'),'clear'); sh.set(qn('w:fill'),hexcolor)
    tcPr.append(sh)

USABLE = 7.66  # 8.5in - 0.42*2 margins

def fix_widths(tbl, fracs):
    """Word only honours column widths when autofit is off and every cell agrees."""
    tbl.autofit = False
    tblPr = tbl._tbl.tblPr
    lay = OxmlElement('w:tblLayout'); lay.set(qn('w:type'),'fixed'); tblPr.append(lay)
    tw = OxmlElement('w:tblW'); tw.set(qn('w:w'), str(int(USABLE*1440))); tw.set(qn('w:type'),'dxa')
    tblPr.append(tw)
    tot = sum(fracs)
    widths = [USABLE*f/tot for f in fracs]
    # rewrite the grid — Word/LibreOffice lay out from w:tblGrid, not from cell widths
    grid = tbl._tbl.find(qn('w:tblGrid'))
    if grid is not None:
        for gc in list(grid): grid.remove(gc)
        for w in widths:
            gc = OxmlElement('w:gridCol'); gc.set(qn('w:w'), str(int(w*1440))); grid.append(gc)
    for row in tbl.rows:
        for ci, c in enumerate(row.cells):
            if ci < len(widths): c.width = Inches(widths[ci])
    return tbl

def borders(tbl):
    tblPr = tbl._tbl.tblPr; b = OxmlElement('w:tblBorders')
    for e in ('top','left','bottom','right','insideH','insideV'):
        x = OxmlElement('w:'+e); x.set(qn('w:val'),'single'); x.set(qn('w:sz'),'4')
        x.set(qn('w:color'),'d9d2bf'); b.append(x)
    tblPr.append(b)

def sz(pt): return Pt(pt*SCALE)

def add_runs(par, node, bold=False, italic=False, color=None, size=8.0):
    """Walk INLINE descendants into runs of the SAME paragraph."""
    for ch in node.children:
        if isinstance(ch, NavigableString):
            t = re.sub(r'\s+',' ', str(ch))
            if t.strip()=='' and not par.runs: continue
            if t=='' : continue
            r = par.add_run(t); r.bold=bold; r.italic=italic
            r.font.size=sz(size)
            if color is not None: r.font.color.rgb=color
        elif isinstance(ch, Tag):
            if ch.name=='br':
                par.add_run().add_break(); continue
            if ch.name in ('script','style','svg'): continue
            nb = bold or ch.name in ('b','strong')
            ni = italic or ch.name in ('i','em')
            nc = color
            cls = ' '.join(ch.get('class') or [])
            if 'by' in cls or 'act' in cls or 'focus' in cls: nc = GOLD
            if ch.name in INLINE:
                add_runs(par, ch, nb, ni, nc, size)
            else:
                add_runs(par, ch, nb, ni, nc, size)

def block(doc, node, size=8.0, bold=False, italic=False, color=None,
          align=None, space_after=2, style=None, prefix=None, keep=False):
    p = doc.add_paragraph()
    pf = p.paragraph_format
    pf.space_before = Pt(0); pf.space_after = Pt(space_after*SCALE)
    pf.line_spacing = 1.0
    if align is not None: p.alignment = align
    if keep: pf.keep_with_next = True
    if prefix:
        r=p.add_run(prefix); r.bold=True; r.font.size=sz(size); r.font.color.rgb=GOLD
    if isinstance(node, str):
        r=p.add_run(node); r.bold=bold; r.italic=italic; r.font.size=sz(size)
        if color is not None: r.font.color.rgb=color
    else:
        add_runs(p, node, bold, italic, color, size)
    if not p.runs:
        # empty block -> drop it
        p._element.getparent().remove(p._element)
        return None
    return p

def flex_text(node, sep='  ·  '):
    """CSS flex rows have no source whitespace — join child blocks explicitly."""
    parts=[]
    for ch in node.find_all(recursive=False):
        t=re.sub(r'\s+',' ',ch.get_text(' ',strip=True))
        if t: parts.append(t)
    if not parts:
        t=re.sub(r'\s+',' ',node.get_text(' ',strip=True))
        return t
    return sep.join(parts)

def make_table(doc, rows, widths=None, header=False, size=7.2, shades=None):
    """rows: list of list of (Tag|str). Returns the docx table."""
    ncol=max(len(r) for r in rows)
    t=doc.add_table(rows=0, cols=ncol); t.alignment=WD_TABLE_ALIGNMENT.CENTER
    borders(t)
    for ri,row in enumerate(rows):
        cells=t.add_row().cells
        for ci in range(ncol):
            c=cells[ci]
            c.paragraphs[0]._element.getparent().remove(c.paragraphs[0]._element)
            val=row[ci] if ci < len(row) else ''
            p=c.add_paragraph(); pf=p.paragraph_format
            pf.space_before=Pt(0); pf.space_after=Pt(0.6*SCALE); pf.line_spacing=1.0
            if isinstance(val,str):
                if val:
                    r=p.add_run(val); r.font.size=sz(size); r.bold=(header and ri==0)
            else:
                add_runs(p,val,bold=(header and ri==0),size=size)
            if shades:
                sh=shades.get((ri,ci))
                if sh: shade(c,sh)
    if widths:
        for ri in range(len(t.rows)):
            for ci,w in enumerate(widths[:ncol]):
                t.rows[ri].cells[ci].width=Inches(w)
    return t

# ---------------------------------------------------------------- FULL
def build_full(soup, doc):
    page=soup.find('div',class_='page')
    tables=0
    for node in page.find_all(recursive=False):
        cls=' '.join(node.get('class') or [])
        if node.name=='div' and 'topbar' in cls:
            block(doc, flex_text(node), size=6.6, color=SOFT, align=WD_ALIGN_PARAGRAPH.CENTER, space_after=3)
        elif node.name=='div' and 'masthead' in cls:
            h1=node.find('h1'); block(doc, h1.get_text(' ',strip=True), size=20, bold=True,
                                      align=WD_ALIGN_PARAGRAPH.CENTER, space_after=1)
            sub=node.find('div',class_='sub'); block(doc, sub.get_text(' ',strip=True), size=7,
                                      color=GOLD, bold=True, align=WD_ALIGN_PARAGRAPH.CENTER, space_after=2)
            dl=node.find('div',class_='dateline'); block(doc, dl, size=8, align=WD_ALIGN_PARAGRAPH.CENTER, space_after=4)
        elif node.name=='div' and 'classline' in cls:
            block(doc, flex_text(node), size=6.8, color=SOFT, bold=True,
                  align=WD_ALIGN_PARAGRAPH.CENTER, space_after=4)
        elif node.name=='h2':
            num=node.find('span',class_='num')
            txt=(num.get_text(strip=True)+'  ' if num else '')+node.get_text(' ',strip=True).replace(
                 num.get_text(strip=True) if num else '','',1).strip()
            block(doc, txt, size=11, bold=True, color=BAR, space_after=3, keep=True)
        elif node.name=='div' and 'callout' in cls:
            tables += emit_callout(doc,node)
        elif node.name=='table':
            emit_html_table(doc,node); tables+=1
        elif node.name=='div' and 'src' in cls:
            block(doc,node,size=6.4,italic=True,color=SOFT,space_after=3)
        elif node.name=='p':
            pcls=cls
            if 'conf' in pcls: block(doc,node,size=6.4,color=SOFT,space_after=3)
            elif 'note' in pcls: block(doc,node,size=7.0,space_after=3)
            else: block(doc,node,size=7.6,space_after=3)
        elif node.name=='h3':
            block(doc,node.get_text(' ',strip=True),size=8.4,bold=True,color=GOLD,space_after=2,keep=True)
        elif node.name=='ul':
            for li in node.find_all('li',recursive=False):
                block(doc,li,size=7.6,space_after=2,prefix='•  ')
        elif node.name=='div' and 'flag' in cls:
            lvl=node.find('div',class_='lvl'); txt=node.find('div',class_='txt') or node.find('div',class_='t')
            t=make_table(doc,[[lvl.get_text(strip=True).upper(), txt]],widths=[0.62,6.6],size=7.4,
                         shades={(0,0):'a4331f'})
            fix_widths(t,[7,93])
            t.rows[0].cells[0].paragraphs[0].runs[0].font.color.rgb=RGBColor(0xff,0xff,0xff)
            t.rows[0].cells[0].paragraphs[0].runs[0].bold=True
            tables+=1
        elif node.name=='div' and 'map-wrap' in cls:
            if os.path.exists('map_20260905.png'):
                p=doc.add_paragraph(); p.alignment=WD_ALIGN_PARAGRAPH.CENTER
                p.paragraph_format.space_after=Pt(2)
                p.add_run().add_picture('map_20260905.png', width=Inches(5.6))
            cap=node.find('div',class_='map-cap')
            if cap: block(doc,cap,size=6.6,italic=True,color=SOFT,space_after=3)
        elif node.name=='div' and 'cols' in cls:
            cells=node.find_all('div',recursive=False)
            rows=[[render_col(c) for c in cells]]
            t=doc.add_table(rows=1,cols=len(cells)); borders(t)
            for ci,c in enumerate(cells):
                cell=t.rows[0].cells[ci]
                cell.paragraphs[0]._element.getparent().remove(cell.paragraphs[0]._element)
                fill_col(cell,c)
                pass
            fix_widths(t,[1,1])
            tables+=1
        elif node.name=='div' and 'craft' in cls:
            tables+=emit_craft(doc,node)
        elif node.name=='div' and 'contacts' in cls:
            cols=node.find_all('div',recursive=False)
            n=max(len(c.find_all('div',class_='r')) for c in cols)
            rows=[]
            for i in range(n):
                row=[]
                for c in cols:
                    rs=c.find_all('div',class_='r')
                    row.append(flex_text(rs[i],sep='   ') if i<len(rs) else '')
                rows.append(row)
            t=make_table(doc,rows,widths=[3.6,3.6],size=7.2); fix_widths(t,[1,1]); tables+=1
        elif node.name=='div' and 'foot' in cls:
            block(doc,node,size=6.6,color=SOFT,align=WD_ALIGN_PARAGRAPH.CENTER,space_after=0)
    return tables

def render_col(c): return c
def fill_col(cell,c):
    for ch in c.find_all(recursive=False):
        if ch.name=='h3':
            p=cell.add_paragraph(); p.paragraph_format.space_after=Pt(1.5*SCALE); p.paragraph_format.line_spacing=1.0
            r=p.add_run(ch.get_text(' ',strip=True)); r.bold=True; r.font.size=sz(7.6); r.font.color.rgb=GOLD
        elif ch.name=='ul':
            for li in ch.find_all('li',recursive=False):
                p=cell.add_paragraph(); p.paragraph_format.space_after=Pt(1.5*SCALE); p.paragraph_format.line_spacing=1.0
                r=p.add_run('•  '); r.font.size=sz(7.2); r.bold=True; r.font.color.rgb=GOLD
                add_runs(p,li,size=7.2)
        elif ch.name=='p':
            p=cell.add_paragraph(); p.paragraph_format.space_after=Pt(1.5*SCALE); p.paragraph_format.line_spacing=1.0
            add_runs(p,ch,size=7.2)

def emit_callout(doc,node):
    t=0
    for ch in node.find_all(recursive=False):
        c=' '.join(ch.get('class') or [])
        if ch.name=='p' and 'lead' in c:
            focus=ch.find('p',class_='focus')
            if focus: focus.extract()
            block(doc,ch,size=8.2,space_after=3)
            if focus: block(doc,focus,size=7.8,color=GOLD,space_after=3)
        elif ch.name=='div' and 'strip' in c:
            bar=ch.find('div',class_='bar')
            if bar: block(doc,bar.get_text(' ',strip=True),size=7,bold=True,color=BAR,space_after=2,keep=True)
            rows=[]
            for r in ch.find_all('div',class_='row'):
                kids=r.find_all('div',recursive=False)
                rows.append([kids[0].get_text(' ',strip=True) if kids else '', kids[1] if len(kids)>1 else ''])
            if rows:
                tt=make_table(doc,rows,widths=[1.0,6.2],size=7.2); fix_widths(tt,[12,88]); t+=1
    return t

def emit_craft(doc,node):
    kick=node.find('span',class_='kick')
    if kick: block(doc,kick.get_text(' ',strip=True).upper(),size=6.6,bold=True,color=GOLD,space_after=2,keep=True)
    h4=node.find('h4')
    if h4: block(doc,h4.get_text(' ',strip=True),size=10,bold=True,color=BAR,space_after=2,keep=True)
    for p in node.find_all('p',recursive=False):
        c=' '.join(p.get('class') or [])
        if 'quote' in c or c=='q': block(doc,p,size=7.8,italic=True,color=SOFT,space_after=3)
        else: block(doc,p,size=7.8,space_after=3)
    return 0

def emit_html_table(doc,tb):
    rows=[]; header=False
    th=tb.find('thead')
    if th:
        header=True
        rows.append([c for c in th.find_all('tr')[0].find_all(['th','td'])])
    body=tb.find('tbody') or tb
    for tr in body.find_all('tr'):
        rows.append([c for c in tr.find_all(['td','th'])])
    ncol=max(len(r) for r in rows)
    fracs=None
    if header:
        got=[]
        for c in rows[0]:
            m=re.search(r'width:\s*([\d.]+)%', c.get('style') or '')
            got.append(float(m.group(1)) if m else None)
        if any(g is not None for g in got):
            miss=[i for i,g in enumerate(got) if g is None]
            rest=max(0.0, 100.0-sum(g for g in got if g))
            for i in miss: got[i]=rest/len(miss) if miss else 1.0
            fracs=[max(g,3.0) for g in got]
    t=doc.add_table(rows=0,cols=ncol); borders(t); t.alignment=WD_TABLE_ALIGNMENT.CENTER
    for ri,row in enumerate(rows):
        cells=t.add_row().cells
        for ci in range(ncol):
            cell=cells[ci]
            cell.paragraphs[0]._element.getparent().remove(cell.paragraphs[0]._element)
            p=cell.add_paragraph(); p.paragraph_format.space_before=Pt(0)
            p.paragraph_format.space_after=Pt(0.6*SCALE); p.paragraph_format.line_spacing=1.0
            if ci<len(row):
                node=row[ci]
                chip=node.find('span',class_=re.compile(r'\bchip\b')) if hasattr(node,'find') else None
                badge=node.find('span',class_=re.compile(r'\bbadge\b')) if hasattr(node,'find') else None
                add_runs(p,node,bold=(header and ri==0),size=7.0)
                key=None
                if chip: key=[x for x in chip.get('class') if x in CHIPBG]
                if badge: key=[x for x in badge.get('class') if x in BADGEBG]
                if key:
                    fill=CHIPBG.get(key[0]) or BADGEBG.get(key[0])
                    if fill:
                        if ncol>=6:
                            col=RGBColor.from_string(fill.upper())
                            for r in p.runs: r.font.color.rgb=col; r.bold=True
                        else:
                            shade(cell,fill)
                            for r in p.runs: r.font.color.rgb=RGBColor(0xff,0xff,0xff); r.bold=True
            if header and ri==0: shade(cell,'211f18')
        if header and ri==0:
            for c in cells:
                for p in c.paragraphs:
                    for r in p.runs: r.font.color.rgb=RGBColor(0xff,0xff,0xff)
    if fracs: fix_widths(t,fracs)
    return t

# ---------------------------------------------------------------- DBS
def build_dbs(soup, doc):
    page=soup.find('div',class_='page'); tables=0
    for node in page.find_all(recursive=False):
        if node.name=='script': continue
        cls=' '.join(node.get('class') or [])
        if 'topbar' in cls:
            block(doc,flex_text(node),size=6.6,color=SOFT,align=WD_ALIGN_PARAGRAPH.CENTER,space_after=3)
        elif 'masthead' in cls:
            block(doc,node.find('h1').get_text(' ',strip=True),size=20,bold=True,align=WD_ALIGN_PARAGRAPH.CENTER,space_after=1)
            block(doc,node.find('div',class_='sub').get_text(' ',strip=True),size=7,color=GOLD,bold=True,
                  align=WD_ALIGN_PARAGRAPH.CENTER,space_after=2)
            block(doc,node.find('div',class_='dateline'),size=8,align=WD_ALIGN_PARAGRAPH.CENTER,space_after=4)
        elif 'classline' in cls:
            block(doc,flex_text(node),size=6.8,color=SOFT,bold=True,align=WD_ALIGN_PARAGRAPH.CENTER,space_after=4)
        elif 'band' in cls:
            cells=node.find_all('div',class_='cell')
            r1=[c.find('div',class_='lab').get_text(' ',strip=True) for c in cells]
            r2=[c.find('span',class_=re.compile(r'\bchip\b')).get_text(' ',strip=True) for c in cells]
            r3=[c.find('div',class_='tr').get_text(' ',strip=True) for c in cells]
            sh={}
            for i,c in enumerate(cells):
                k=[x for x in c.find('span',class_=re.compile(r'\bchip\b')).get('class') if x in CHIPBG]
                if k: sh[(1,i)]=CHIPBG[k[0]]
            t=make_table(doc,[r1,r2,r3],widths=[1.44]*5,size=6.6,shades=sh); fix_widths(t,[1]*5)
            for i in range(5):
                for p in t.rows[1].cells[i].paragraphs:
                    for r in p.runs: r.font.color.rgb=RGBColor(0xff,0xff,0xff); r.bold=True
            tables+=1
        elif 'lead' in cls:
            focus=node.find('div',class_='focus')
            if focus: focus.extract()
            block(doc,node,size=8.0,space_after=3)
            if focus: block(doc,focus,size=7.6,color=GOLD,space_after=3)
        elif node.name=='h3':
            main=node.get_text(' ',strip=True); em=node.find('em')
            if em: main=main.replace(em.get_text(' ',strip=True),'').strip()
            block(doc,main.upper(),size=8.4,bold=True,color=BAR,space_after=2,keep=True)
        elif 'glance' in cls:
            left=node.find('div',class_='chart'); right=node.find('div',class_='out')
            t=doc.add_table(rows=1,cols=2); borders(t)
            c0=t.rows[0].cells[0]; c0.paragraphs[0]._element.getparent().remove(c0.paragraphs[0]._element)
            p=c0.add_paragraph(); p.alignment=WD_ALIGN_PARAGRAPH.CENTER
            if os.path.exists('chart_20260905.png'):
                p.add_run().add_picture('chart_20260905.png', width=Inches(4.25))
            c1=t.rows[0].cells[1]; c1.paragraphs[0]._element.getparent().remove(c1.paragraphs[0]._element)
            for d in right.find_all('div',class_='d'):
                pw=c1.add_paragraph(); pw.paragraph_format.space_after=Pt(0.5*SCALE); pw.paragraph_format.line_spacing=1.0
                r=pw.add_run(d.find('div',class_='when').get_text(' ',strip=True).upper())
                r.bold=True; r.font.size=sz(6.4); r.font.color.rgb=RED
                pp=c1.add_paragraph(); pp.paragraph_format.space_after=Pt(2*SCALE); pp.paragraph_format.line_spacing=1.0
                add_runs(pp,d.find('p'),size=7.0)
            fix_widths(t,[57,43])
            tables+=1
        elif cls.strip()=='dv':
            divs=node.find_all('div',recursive=False)
            rows=[]
            for i in range(0,len(divs),4):
                chunk=divs[i:i+4]
                rows.append([d.find('div',class_='n').get_text(' ',strip=True) for d in chunk])
                rows.append([d.find('p') for d in chunk])
            t=make_table(doc,rows,widths=[1.8]*4,size=6.9); fix_widths(t,[1]*4)
            for ri in range(0,len(t.rows),2):
                for c in t.rows[ri].cells:
                    for p in c.paragraphs:
                        for r in p.runs: r.bold=True; r.font.size=sz(6.4)
            tables+=1
        elif 'wi' in cls.split():
            h=node.find('div',class_='h')
            if h: block(doc,flex_text(h,sep='   —   '),size=6.8,bold=True,color=BAR,space_after=2,keep=True)
            rows=[]; sh={}
            for i,r in enumerate(node.find_all('div',class_='r')):
                ref=r.find('div',class_='ref'); badge=r.find('div',class_=re.compile(r'\bbadge\b')); tx=r.find('div',class_='tx')
                rows.append([ref.get_text(strip=True), badge.get_text(' ',strip=True), tx])
                k=[x for x in badge.get('class') if x in BADGEBG]
                if k: sh[(i,1)]=BADGEBG[k[0]]
            t=make_table(doc,rows,widths=[0.45,0.95,5.85],size=7.1,shades=sh)
            fix_widths(t,[5.5,10.5,84.0])
            for (ri,ci),_ in sh.items():
                for p in t.rows[ri].cells[ci].paragraphs:
                    for r in p.runs: r.font.color.rgb=RGBColor(0xff,0xff,0xff); r.bold=True
            tables+=1
        elif 'two' in cls.split():
            cols=node.find_all('div',recursive=False)
            t=doc.add_table(rows=1,cols=2); borders(t)
            for ci,c in enumerate(cols):
                cell=t.rows[0].cells[ci]
                cell.paragraphs[0]._element.getparent().remove(cell.paragraphs[0]._element)
                fill_two(cell,c)
                pass
            fix_widths(t,[55,45])
            tables+=1
        elif 'craft' in cls:
            kick=node.find('span',class_='kick')
            block(doc,kick.get_text(' ',strip=True).upper(),size=6.6,bold=True,color=GOLD,space_after=2,keep=True)
            for p in node.find_all('p',recursive=False):
                c=' '.join(p.get('class') or [])
                block(doc,p,size=7.6,italic=('q' in c.split()),color=(SOFT if 'q' in c.split() else None),space_after=3)
        elif 'ready' in cls.split():
            kick=node.find('div',class_='kick')
            block(doc,flex_text(kick,sep='   —   ').upper(),size=6.6,bold=True,color=BAR,space_after=2,keep=True)
            g=node.find('div',class_='rgrid'); divs=g.find_all('div',recursive=False)
            rows=[[d.find('div',class_='n').get_text(' ',strip=True) for d in divs],
                  [d.find('p') for d in divs]]
            t=make_table(doc,rows,widths=[2.4]*3,size=7.1); fix_widths(t,[1]*3)
            for c in t.rows[0].cells:
                for p in c.paragraphs:
                    for r in p.runs: r.bold=True; r.font.size=sz(6.4); r.font.color.rgb=SOFT
            tables+=1
        elif 'strip' in cls.split():
            cs=node.find_all('div',class_='c')
            labs=[]; vals=[]
            for c in cs:
                b=c.find('b'); v=b.get_text(' ',strip=True) if b else ''
                if b: b.extract()
                labs.append(c.get_text(' ',strip=True)); vals.append(v)
            t=make_table(doc,[labs,vals],widths=[1.2]*len(cs),size=6.6); fix_widths(t,[1]*len(cs))
            for c in t.rows[1].cells:
                for p in c.paragraphs:
                    for r in p.runs: r.bold=True; r.font.size=sz(7.4)
            tables+=1
        elif 'foot' in cls.split():
            block(doc,node,size=6.4,color=SOFT,align=WD_ALIGN_PARAGRAPH.CENTER,space_after=0)
    return tables

def fill_two(cell,c):
    for ch in c.find_all(recursive=False):
        if ch.name=='h3':
            p=cell.add_paragraph(); p.paragraph_format.space_after=Pt(1.5*SCALE); p.paragraph_format.line_spacing=1.0
            r=p.add_run(ch.get_text(' ',strip=True).upper()); r.bold=True; r.font.size=sz(7.4); r.font.color.rgb=BAR
        elif ch.name=='ul':
            for li in ch.find_all('li',recursive=False):
                p=cell.add_paragraph(); p.paragraph_format.space_after=Pt(1.5*SCALE); p.paragraph_format.line_spacing=1.0
                r=p.add_run('•  '); r.bold=True; r.font.size=sz(7.0); r.font.color.rgb=GOLD
                add_runs(p,li,size=7.0)
        elif ch.name=='div' and 'flag' in ' '.join(ch.get('class') or []):
            lvl=ch.find('div',class_='lvl'); t=ch.find('div',class_='t')
            p=cell.add_paragraph(); p.paragraph_format.space_after=Pt(1.5*SCALE); p.paragraph_format.line_spacing=1.0
            r=p.add_run(lvl.get_text(strip=True).upper()+'  '); r.bold=True; r.font.size=sz(7.4); r.font.color.rgb=RED
            add_runs(p,t,size=7.0)
        elif ch.name=='div' and 'q' in (ch.get('class') or []):
            p=cell.add_paragraph(); p.paragraph_format.space_after=Pt(0.5*SCALE); p.paragraph_format.line_spacing=1.0
            r=p.add_run(ch.get_text(' ',strip=True)); r.bold=True; r.italic=True; r.font.size=sz(7.0); r.font.color.rgb=GOLD
        elif ch.name=='div' and 'a' in (ch.get('class') or []):
            p=cell.add_paragraph(); p.paragraph_format.space_after=Pt(2*SCALE); p.paragraph_format.line_spacing=1.0
            add_runs(p,ch,size=7.0)

def build(kind, scale):
    global SCALE; SCALE=scale
    src=f'DaybreakBrief_{kind}_20260905.html'
    soup=BeautifulSoup(open(src,encoding='utf-8').read(),'lxml')
    for s in soup.find_all('script'): s.decompose()
    doc=Document()
    st=doc.styles['Normal']; st.font.name='Calibri'; st.font.size=sz(7.8)
    st.element.rPr.rFonts.set(qn('w:eastAsia'),'Calibri')
    for s in doc.sections:
        s.page_width=Inches(8.5); s.page_height=Inches(11)
        s.top_margin=Inches(0.30); s.bottom_margin=Inches(0.26)
        s.left_margin=Inches(0.42); s.right_margin=Inches(0.42)
    n = build_full(soup,doc) if kind=='FULL' else build_dbs(soup,doc)
    out=f'DaybreakBrief_{kind}_20260905.docx'
    doc.save(out)
    return out,n

if __name__=='__main__':
    kind=sys.argv[1]; scale=float(sys.argv[2]) if len(sys.argv)>2 else 1.0
    out,n=build(kind,scale)
    print(out, 'tables:', n, 'scale:', scale, 'bytes:', os.path.getsize(out))
